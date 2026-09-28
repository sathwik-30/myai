import csv
import json
from pathlib import Path
from typing import Dict, List

# Temporary resources are files supplied for a current exam/project/task.
RESOURCE_ROOT = Path(__file__).resolve().parents[1] / "memory" / "temporary" / "resources"
SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".pdf", ".pptx", ".docx", ".csv", ".json"
}
IGNORED_FILENAMES = {"README.md"}


def _read_pdf(path: Path) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation
    presentation = Presentation(str(path))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    from docx import Document
    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".pptx":
        return _read_pptx(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".csv":
        with path.open("r", encoding="utf-8", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    if suffix == ".json":
        return json.dumps(json.loads(path.read_text(encoding="utf-8")), indent=2, ensure_ascii=False)
    return ""


def load_resources() -> List[Dict[str, str]]:
    RESOURCE_ROOT.mkdir(parents=True, exist_ok=True)
    resources = []
    ignored = {name.lower() for name in IGNORED_FILENAMES}

    for path in sorted(RESOURCE_ROOT.rglob("*")):
        if (
            not path.is_file()
            or path.suffix.lower() not in SUPPORTED_EXTENSIONS
            or path.name.lower() in ignored
        ):
            continue

        try:
            text = _read_file(path)
        except (OSError, UnicodeDecodeError, ValueError, ImportError):
            continue

        if text.strip():
            resources.append({
                "name": path.name,
                "path": str(path.relative_to(RESOURCE_ROOT)),
                "text": text.strip(),
            })

    return resources


def _relevant_excerpt(text: str, query_words: set[str], limit: int = 1800) -> str:
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [line.strip() for line in text.splitlines() if line.strip()]

    scored = []
    for index, paragraph in enumerate(paragraphs):
        words = {word.lower() for word in paragraph.split() if len(word) > 2}
        overlap = len(query_words & words)
        if overlap:
            scored.append((overlap, -index, paragraph))

    if not scored:
        return text[:limit]

    scored.sort(reverse=True)
    selected = [item[2] for item in scored[:3]]
    return "\n\n".join(selected)[:limit]


def search_resources(query: str, limit: int = 3) -> List[Dict[str, str]]:
    resources = load_resources()
    if not resources:
        return []

    query_words = {word.lower() for word in query.split() if len(word) > 2}
    if not query_words:
        return []

    scored = []
    for item in resources:
        words = {word.lower() for word in item["text"].split() if len(word) > 2}
        overlap = len(query_words & words)
        if overlap:
            scored.append((overlap / max(len(query_words), 1), item))

    scored.sort(key=lambda pair: pair[0], reverse=True)

    results = []
    for score, item in scored[:limit]:
        if score < 0.20:
            continue
        result = dict(item)
        result["text"] = _relevant_excerpt(item["text"], query_words)
        result["score"] = round(score, 4)
        results.append(result)

    return results
